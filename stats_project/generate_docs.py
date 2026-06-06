"""
一键生成 Word 报告 + PowerPoint 答辩PPT
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PptxInches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor as PptxRGBColor
import os

# ==================== 1. 生成 Word 报告 ====================
print("正在生成 Word 报告...")
doc = Document()

# 设置默认中文字体
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

# 标题
title = doc.add_heading('买断制独立游戏销量的统计分布特征与参数估计', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.runs[0]
run.font.size = Pt(18)
run.font.bold = True
run.font.color.rgb = RGBColor(0, 0, 128)

# 副标题
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('概率统计课程项目报告\n基于 MLE、Bootstrap 与回归分析')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(80, 80, 80)

doc.add_paragraph()

# 摘要
h = doc.add_heading('一、摘要', level=1)
doc.add_paragraph(
    '本项目以买断制独立游戏市场为研究对象，基于模拟生成的 500 款独立游戏样本数据，'
    '系统分析了游戏评论数（销量代理变量）的统计分布特征。主要工作包括：描述统计、'
    'MLE 最大似然估计、K-S 拟合优度检验、Bootstrap 置信区间估计，以及价格与好评率的一元线性回归分析。'
    '核心结论：独立游戏评论数服从对数正态分布，呈现典型的"长尾效应"；价格与好评率之间无显著线性关系。'
)

# 数据说明
doc.add_heading('二、数据来源与变量', level=1)
doc.add_paragraph(
    '样本量 n = 500，为基于 Steam 平台独立游戏真实统计特征生成的模拟数据集。'
    '核心变量 review_count（评论数）作为销量的代理指标。'
)

# 描述统计
doc.add_heading('三、描述统计', level=1)
doc.add_paragraph('对评论数进行描述统计，结果如下：')
table = doc.add_table(rows=6, cols=3)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = '统计量'
hdr[1].text = '数值'
hdr[2].text = '说明'
rows = [
    ('样本量 n', '500', '—'),
    ('均值', '14003.26', '被极端高值拉高'),
    ('中位数', '1740.50', '半数游戏低于此值'),
    ('偏度 (Skewness)', '3.9576', '显著右偏'),
    ('峰度 (Kurtosis)', '16.6333', '尖峰厚尾'),
]
for i, (name, val, desc) in enumerate(rows, 1):
    r = table.rows[i].cells
    r[0].text = name
    r[1].text = val
    r[2].text = desc

doc.add_paragraph()
doc.add_paragraph(
    '从统计量可见，评论数呈严重右偏分布（偏度>3），且峰度远大于0，说明分布比正态分布更集中于低值区，'
    '同时右侧存在厚尾——即少数"爆款"游戏拥有极高的评论数。'
)
doc.add_picture('output/figures/01_histogram.png', width=Inches(5.5))
last_paragraph = doc.paragraphs[-1]
last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

# MLE
doc.add_heading('四、MLE 分布拟合', level=1)
doc.add_paragraph(
    '假设评论数 X 服从对数正态分布 LogNormal(μ, σ²)。'
    '对数变换后 ln(X) ~ N(μ, σ²)，MLE 估计具有解析解：'
)
doc.add_paragraph('μ̂ = (1/n) Σ ln(xi) = 7.6115')
doc.add_paragraph('σ̂² = (1/n) Σ (ln(xi) − μ̂)² = 4.6338')
doc.add_paragraph('σ̂ = 2.1526')
doc.add_paragraph(
    '下图展示了直方图与拟合密度曲线的对比，红色的对数正态分布拟合效果明显优于绿色的伽马分布。'
)
doc.add_picture('output/figures/02_mle_fitting.png', width=Inches(5.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# K-S检验
doc.add_heading('五、K-S 拟合优度检验', level=1)
doc.add_paragraph('原假设 H0：数据服从对数正态分布。')
doc.add_paragraph('备择假设 H1：数据不服从对数正态分布。')
doc.add_paragraph('检验结果：K-S 统计量 = 0.0300，p-value = 0.7483。')
doc.add_paragraph(
    '结论：在 5% 显著性水平下，p > 0.05，不能拒绝 H0，即样本数据与对数正态分布拟合良好。'
)
doc.add_picture('output/figures/03_qqplot.png', width=Inches(5.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Bootstrap
doc.add_heading('六、Bootstrap 置信区间', level=1)
doc.add_paragraph(
    '由于原始数据高度右偏，传统基于正态假设的置信区间可能不准确。'
    '采用 Bootstrap 自助法，有放回重抽样 B = 5000 次，得到样本均值的 95% 置信区间：'
)
doc.add_paragraph('百分位数 Bootstrap 95% CI: [11182.27, 17114.56]')
doc.add_paragraph('正态近似 95% CI: [11032.75, 16973.77]')
doc.add_paragraph(
    '两种方法结果接近，说明样本量 n=500 足够大，中心极限定理已发挥作用。'
    'Bootstrap 均值分布近似正态，直观验证了 CLT 的强大性。'
)
doc.add_picture('output/figures/04_bootstrap.png', width=Inches(5.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# 回归分析
doc.add_heading('七、一元线性回归（拓展）', level=1)
doc.add_paragraph('研究问题：独立游戏的价格是否影响用户好评率？')
doc.add_paragraph('模型：好评率 = β0 + β1 × 价格 + ε')
doc.add_paragraph('OLS 估计结果：')
doc.add_paragraph('回归方程: 好评率 = 0.7911 + 0.000100 × 价格')
doc.add_paragraph('R² = 0.0007，价格系数 p-value = 0.5557')
doc.add_paragraph(
    '结论：价格对好评率无显著影响（p ≥ 0.05），R² 极低，说明好评率主要由游戏品质驱动，而非定价策略。'
)
doc.add_picture('output/figures/05_regression.png', width=Inches(5.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_picture('output/figures/06_residuals.png', width=Inches(5.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# 结论
doc.add_heading('八、结论', level=1)
doc.add_paragraph(
    '1. 独立游戏评论数服从对数正态分布，市场呈典型的"长尾效应"与"赢家通吃"格局。\n'
    '2. MLE 能有效估计分布参数，K-S 检验验证拟合良好（p = 0.7483）。\n'
    '3. Bootstrap 提供稳健的置信区间估计，且不依赖严格的正态假设。\n'
    '4. 价格对好评率无显著线性影响，独立游戏的口碑主要由品质驱动。'
)

doc.save('output/统计课程项目报告.docx')
print("✅ Word 报告已生成: output/统计课程项目报告.docx")

# ==================== 2. 生成 PPT ====================
print("正在生成 PPT...")
prs = Presentation()
prs.slide_width = PptxInches(13.333)
prs.slide_height = PptxInches(7.5)

def add_title_slide(prs, title_text, subtitle_text):
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)
    title_box = slide.shapes.add_textbox(PptxInches(1), PptxInches(2), PptxInches(11.333), PptxInches(1.5))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = PptxRGBColor(0, 51, 102)
    p.alignment = PP_ALIGN.CENTER

    sub_box = slide.shapes.add_textbox(PptxInches(1), PptxInches(3.8), PptxInches(11.333), PptxInches(1))
    tf2 = sub_box.text_frame
    tf2.text = subtitle_text
    p2 = tf2.paragraphs[0]
    p2.font.size = Pt(18)
    p2.font.color.rgb = PptxRGBColor(80, 80, 80)
    p2.alignment = PP_ALIGN.CENTER
    return slide

def add_content_slide(prs, title, bullets, img_path=None, img_left=PptxInches(7), img_top=PptxInches(1.5), img_width=PptxInches(5.5)):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # 标题
    title_box = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(0.3), PptxInches(12), PptxInches(0.8))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = PptxRGBColor(0, 51, 102)

    # 内容
    content_box = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(1.2), PptxInches(11.5), PptxInches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, text in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(18)
        p.space_after = Pt(12)
        p.level = 0

    if img_path and os.path.exists(img_path):
        slide.shapes.add_picture(img_path, img_left, img_top, width=img_width)

    return slide

# 第1页：标题
add_title_slide(prs,
    '买断制独立游戏销量的统计分布特征与参数估计',
    '基于 MLE、Bootstrap 与回归分析的概率统计课程项目')

# 第2页：背景
add_content_slide(prs, '研究背景与核心问题', [
    '• 独立游戏市场蓬勃发展，但高度分化',
    '• 少数"爆款" vs 大量"长尾"作品',
    '• 核心问题：',
    '   1. 独立游戏销量服从什么分布？',
    '   2. 如何用 MLE 估计分布参数？',
    '   3. 游戏价格是否影响用户好评率？'
])

# 第3页：数据
add_content_slide(prs, '数据来源与变量', [
    '• 样本量：n = 500 款买断制独立游戏（模拟数据集）',
    '• 核心变量：review_count（评论数）→ 销量代理指标',
    '• 其他变量：价格、游戏时长、好评率、中文支持、标签数、发行年份',
    '• 选择理由：Steam 不公开销量，但评论数与销量比例稳定（约 1:30 ~ 1:50）'
])

# 第4页：描述统计
add_content_slide(prs, '描述统计：评论数呈显著右偏', [
    '• 均值 = 14003.26  >>  中位数 = 1740.50',
    '• 偏度 (Skewness) = 3.96  > 0，显著右偏',
    '• 峰度 (Kurtosis) = 16.63，尖峰厚尾',
    '• 对数变换后接近正态分布 → 暗示对数正态分布'
], 'output/figures/01_histogram.png', PptxInches(6.5), PptxInches(1.3), PptxInches(6))

# 第5页：MLE（核心）
add_content_slide(prs, 'MLE 参数估计：对数正态分布', [
    '• 假设：X ~ LogNormal(μ, σ²)，ln(X) ~ N(μ, σ²)',
    '• MLE 解析解：',
    '   μ̂ = (1/n)Σln(xi) = 7.61',
    '   σ̂ = 2.15',
    '• 拟合效果：对数正态分布 > 伽马分布',
    '• 优点：大样本下具有一致性、有效性（达到 CR 下界）'
], 'output/figures/02_mle_fitting.png', PptxInches(6.5), PptxInches(1.3), PptxInches(6))

# 第6页：K-S检验
add_content_slide(prs, '拟合优度检验：K-S 检验 + Q-Q 图', [
    '• H0：数据服从对数正态分布',
    '• H1：数据不服从对数正态分布',
    '• K-S 统计量 = 0.0300，p-value = 0.7483',
    '• 结论：p > 0.05，不能拒绝 H0，拟合良好',
    '• Q-Q 图：样本分位数与理论分位数基本落在参考线上'
], 'output/figures/03_qqplot.png', PptxInches(6.5), PptxInches(1.3), PptxInches(6))

# 第7页：Bootstrap
add_content_slide(prs, 'Bootstrap 自助法：稳健区间估计', [
    '• 传统方法依赖正态假设，对偏态分布不准确',
    '• Bootstrap：有放回重抽样 B = 5000 次',
    '• 百分位数法 95% CI: [11182, 17115]',
    '• 正态近似 95% CI: [11033, 16974]',
    '• Bootstrap 均值分布近似正态 → 验证中心极限定理（CLT）'
], 'output/figures/04_bootstrap.png', PptxInches(6.5), PptxInches(1.3), PptxInches(6))

# 第8页：回归
add_content_slide(prs, '拓展分析：价格对好评率的影响', [
    '• 模型：好评率 = β0 + β1 × 价格 + ε',
    '• OLS 估计：好评率 = 0.7911 + 0.000100 × 价格',
    '• R² = 0.001，解释力极弱',
    '• 价格系数 p-value = 0.556 > 0.05，无显著影响',
    '• 结论：玩家更在意游戏品质，而非价格'
], 'output/figures/05_regression.png', PptxInches(6.5), PptxInches(1.3), PptxInches(6))

# 第9页：结论
add_content_slide(prs, '主要结论', [
    '1. 独立游戏评论数服从对数正态分布，市场呈"长尾效应"',
    '2. MLE 能有效估计分布参数，K-S 检验验证拟合良好',
    '3. Bootstrap 提供稳健的置信区间估计',
    '4. 价格对好评率无显著影响，口碑由品质驱动',
    '',
    '对开发者的启示：',
    '• 做好"大部分游戏不会爆火"的市场预期',
    '• 不必过度担心定价影响口碑，专注打磨品质'
])

# 第10页：致谢
add_title_slide(prs, '谢谢聆听！', '欢迎老师提问')

prs.save('output/答辩PPT.pptx')
print("✅ PPT 已生成: output/答辩PPT.pptx")

print("\n🎉 全部完成！三件大事都搞定了：")
print("   1. 代码运行完毕，6张图表已生成")
print("   2. Word 报告: output/统计课程项目报告.docx")
print("   3. PPT 文件: output/答辩PPT.pptx")
